from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook

FIELDS = ("invoice_number", "amount", "tax", "total", "seller_name")
EXCEL_HEADERS = {
    "invoice_number": "发票号码",
    "amount": "金额",
    "tax": "税额",
    "total": "价税合计金额",
    "seller_name": "销售方公司名称",
}


def normalize(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().replace(",", "").replace("￥", "").replace("¥", "")
    try:
        return f"{float(text):.2f}"
    except ValueError:
        return text


def evaluate(ground_truth: Path, workbook_path: Path) -> dict[str, tuple[int, int]]:
    with ground_truth.open("r", encoding="utf-8-sig", newline="") as stream:
        expected = {
            (row["filename"], str(row.get("page_number", "")).strip()): row
            for row in csv.DictReader(stream)
            if row.get("filename")
        }
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["发票明细"]
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    columns = {str(name): index for index, name in enumerate(header) if name is not None}
    filename_column = columns["原始文件名"]
    page_column = columns.get("PDF页码")
    field_columns = {field: columns[label] for field, label in EXCEL_HEADERS.items()}
    predicted = {
        (
            str(row[filename_column]),
            str(row[page_column]).strip() if page_column is not None and row[page_column] is not None else "",
        ): row
        for row in sheet.iter_rows(min_row=2, values_only=True)
    }
    results: dict[str, tuple[int, int]] = {}
    for field in FIELDS:
        correct = total = 0
        for row_key, expected_row in expected.items():
            if row_key not in predicted:
                continue
            total += 1
            actual = normalize(predicted[row_key][field_columns[field]])
            wanted = normalize(expected_row.get(field, ""))
            correct += actual == wanted
        results[field] = (correct, total)
    workbook.close()
    return results


def main() -> None:
    parser = argparse.ArgumentParser(description="按人工标注 CSV 统计各字段精确匹配率")
    parser.add_argument("ground_truth", type=Path)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    for field, (correct, total) in evaluate(args.ground_truth, args.workbook).items():
        rate = correct / total * 100 if total else 0
        print(f"{field}: {correct}/{total} = {rate:.2f}%")


if __name__ == "__main__":
    main()
