from __future__ import annotations

from collections import Counter
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import InvoiceRecord

DETAIL_HEADERS = [
    "序号", "发票号码", "金额", "税额", "价税合计金额", "销售方公司名称",
    "原始文件名", "PDF页码", "发票代码", "开票日期", "识别状态", "校验结果",
    "图片路径", "识别来源", "本地初次识别值", "识别失败原因",
]


def export_excel(records: list[InvoiceRecord], output_path: str | Path) -> Path:
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    detail = workbook.active
    detail.title = "发票明细"
    exceptions = workbook.create_sheet("异常复核")
    statistics = workbook.create_sheet("处理统计")
    ocr_content = workbook.create_sheet("OCR原始内容")

    _write_table(detail, DETAIL_HEADERS, [_record_row(item) for item in records])
    exception_records = [item for item in records if item.status != "成功"]
    _write_table(exceptions, DETAIL_HEADERS, [_record_row(item) for item in exception_records])
    ocr_rows = []
    for record in records:
        for sequence, line in enumerate(record.raw_ocr, start=1):
            ocr_rows.append([
                record.index,
                record.original_filename,
                line.get("page_number", record.source_page),
                sequence,
                line.get("text", ""),
                json.dumps(line.get("box", []), ensure_ascii=False, separators=(",", ":")),
                json.dumps(line.get("layout_box", []), ensure_ascii=False, separators=(",", ":")),
                line.get("score"),
                line.get("source", "local OCR"),
                line.get("coordinate_unit", "pixel"),
            ])
    _write_table(
        ocr_content,
        ["序号", "原始文件名", "PDF页码", "OCR行序", "OCR原文", "原图坐标框", "校正版面坐标框", "OCR行置信度", "文字来源", "坐标单位"],
        ocr_rows,
    )

    counts = Counter(item.status for item in records)
    _write_table(
        statistics,
        ["统计项", "数量/内容"],
        [
            ["发票记录数", len(records)],
            ["成功", counts.get("成功", 0)],
            ["待复核", counts.get("待复核", 0)],
            ["失败", counts.get("失败", 0)],
            ["本地识别", sum(item.recognition_source.startswith("local") for item in records)],
            ["经用户确认的云端识别", sum(item.recognition_source.startswith("cloud") for item in records)],
            ["备注", "空字段保持为空；金额为票面识别值，不以计算值覆盖。"],
        ],
    )
    workbook.save(destination)
    return destination


def _record_row(record: InvoiceRecord) -> list[object]:
    return [
        record.index,
        record.invoice_number or "",
        float(record.amount) if record.amount is not None else None,
        float(record.tax) if record.tax is not None else None,
        float(record.total) if record.total is not None else None,
        record.seller_name or "",
        record.original_filename,
        ",".join(str(page) for page in record.source_pages) or (str(record.source_page) if record.source_page else ""),
        record.invoice_code or "",
        record.invoice_date or "",
        record.status,
        "；".join(record.validation_result),
        record.file_path,
        record.recognition_source,
        json.dumps(record.original_extracted_fields, ensure_ascii=False, separators=(",", ":")),
        record.error,
    ]


def _write_table(sheet: object, headers: list[str], rows: list[list[object]]) -> None:
    sheet.append(headers)  # type: ignore[attr-defined]
    for row in rows:
        sheet.append(row)  # type: ignore[attr-defined]
        for cell in sheet[sheet.max_row]:  # type: ignore[attr-defined]
            if isinstance(cell.value, str) and cell.value.startswith("="):
                # OCR text and filenames are untrusted input, never formulas.
                cell.data_type = "s"
    sheet.freeze_panes = "A2"  # type: ignore[attr-defined]
    sheet.auto_filter.ref = sheet.dimensions  # type: ignore[attr-defined]
    for cell in sheet[1]:  # type: ignore[index]
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2463A6")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):  # type: ignore[attr-defined]
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in (3, 4, 5):
            if column <= len(row) and isinstance(row[column - 1].value, (float, int)):
                row[column - 1].number_format = "0.00"
    for column in sheet.columns:  # type: ignore[attr-defined]
        letter = get_column_letter(column[0].column)
        max_length = max((len(str(cell.value or "")) for cell in column), default=0)
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 48)  # type: ignore[attr-defined]
    sheet.row_dimensions[1].height = 28  # type: ignore[attr-defined]
    if sheet.title in {"发票明细", "异常复核"}:  # type: ignore[attr-defined]
        for row in range(2, sheet.max_row + 1):  # type: ignore[attr-defined]
            path_cell = sheet.cell(row=row, column=13)  # type: ignore[attr-defined]
            path_value = path_cell.value
            if path_value:
                try:
                    path_cell.hyperlink = Path(path_value).resolve().as_uri()
                    path_cell.style = "Hyperlink"
                except (OSError, ValueError):
                    pass
